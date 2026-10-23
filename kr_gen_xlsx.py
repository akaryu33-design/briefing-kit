"""한국증시 브리핑 엑셀 생성기 (데일리·위클리 공용, 미국증시 gen_xlsx.py 양식 기반).
사용: python kr_gen_xlsx.py data.json out.xlsx
미국판 대비 추가: 투자자별 수급 시트, (위클리) 일별 경로 시트, 국고채 커브·원화·수급 요약 수식."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding="utf-8")); out = sys.argv[2]
FONT = "Malgun Gothic"
BLUE = Font(name=FONT, color="0000FF", size=10); BLACK = Font(name=FONT, size=10)
GREEN = Font(name=FONT, color="008000", size=10)
BOLD = Font(name=FONT, bold=True, size=10); SUB = Font(name=FONT, italic=True, size=9, color="595959")
HF = Font(name=FONT, bold=True, color="FFFFFF", size=10); HFILL = PatternFill("solid", fgColor="1F3864")
TITLE = Font(name=FONT, bold=True, size=14)
b = Side(style="thin", color="BFBFBF"); BD = Border(left=b, right=b, top=b, bottom=b)
WRAP = Alignment(wrap_text=True, vertical="top")
PCT = '+0.00%;-0.00%;0.00%'; BP = '+0.0"bp";-0.0"bp";0.0"bp"'; EOK = '+#,##0;-#,##0;0'
TITLE_TXT = d.get("title", "한국증시 데일리 브리핑")
BASIS = f"기준: {d['date']} {d.get('basis_label', '정규장 마감 기준')} / 작성 {d['written']} {d.get('written_kst', '')} (한국시각)"
CHG = d.get("chg_label", "등락률")
SRC = "Source: " + ", ".join(s[0] for s in d["sources"][:5]) + " 외 (Sources 시트 참조)"
wb = Workbook()


def sheet(title, name, headers, widths, rows, fmts=None, note=SRC):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE; ws["A2"] = BASIS; ws["A2"].font = SUB
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h); c.font = HF; c.fill = HFILL; c.border = BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, 4):
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=r, column=ci, value=v); c.border = BD; c.alignment = WRAP
            num = isinstance(v, (int, float)) and not isinstance(v, bool)
            formula = isinstance(v, str) and v.startswith("=")
            c.font = BLUE if num else BLACK
            if fmts and ci in fmts and (num or formula):
                c.number_format = fmts[ci]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    end = 4 + len(rows)
    if note: ws.cell(row=end, column=1, value=note).font = SUB
    return ws, end


# 1) 지수·매크로: [name, level, pct, bp, note, is_yield, level_fmt]
idx = d["indices"]
rows = [[x[0], x[1], x[2], x[3], x[4]] for x in idx]
ws, end = sheet("지수·매크로", "지수·매크로", ["항목", "종가/수준", CHG, "금리변동", "비고"], [30, 14, 11, 11, 52], rows,
                fmts={2: '#,##0.00', 3: PCT, 4: BP})
row_of = {}
for i, x in enumerate(idx, 4):
    row_of[x[0]] = i
    if len(x) > 5 and x[5]: ws.cell(row=i, column=2).number_format = '0.000%'
    if len(x) > 6 and x[6]: ws.cell(row=i, column=2).number_format = x[6]
br = end + 2
if d.get("breadth"):
    ws.cell(row=br, column=1, value=d.get("breadth_label", "KOSPI 등락 종목 수")).font = BOLD
    ws.cell(row=br, column=2, value="상승").font = BOLD; ws.cell(row=br, column=3, value="하락").font = BOLD
    ws.cell(row=br + 1, column=1, value="종목 수").font = BLACK
    ws.cell(row=br + 1, column=2, value=d["breadth"][0]).font = BLUE; ws.cell(row=br + 1, column=3, value=d["breadth"][1]).font = BLUE


def ref(name, col):
    return f"'지수·매크로'!{col}{row_of[name]}"


# 2) 투자자별 수급 (억원) — 행: 투자자, 열: 시장
fl = d["flows"]; markets = list(fl["markets"])
investors = fl["investors"]
rows = [[inv] + [fl["data"][m].get(inv) for m in markets] for inv in investors]
n = len(rows)
sum_row = [f"{n}주체 합 (잔차 = 미집계 주체·반올림)"] + [f"=SUM({get_column_letter(2 + j)}4:{get_column_letter(2 + j)}{3 + n})" for j in range(len(markets))]
ws, end = sheet(fl.get("title", "투자자별 순매수 (억원)"), "수급", ["투자자"] + markets, [30] + [16] * len(markets),
                rows + [sum_row], fmts={j: EOK for j in range(2, 2 + len(markets))}, note=fl.get("note", SRC))
flow_row = {inv: 4 + k for k, inv in enumerate(investors)}
r = end + 2
for extra in fl.get("extra", []):  # [라벨, 값, 비고]
    ws.cell(row=r, column=1, value=extra[0]).font = BOLD
    c = ws.cell(row=r, column=2, value=extra[1]); c.font = BLUE; c.number_format = EOK
    ws.cell(row=r, column=3, value=extra[2]).font = SUB
    r += 1


def fref(inv, market):
    col = get_column_letter(2 + markets.index(market))
    return f"'수급'!{col}{flow_row[inv]}"


# 3) (위클리) 일별 경로
if d.get("path"):
    p = d["path"]
    rows = [[x[0], x[1], x[2], x[3], x[4]] for x in p["rows"]]
    sheet(p.get("title", "일별 경로"), "일별경로", p["headers"], [16, 12, 11, 14, 60], rows,
          fmts={2: '#,##0.00', 3: PCT, 4: EOK}, note=p.get("note", SRC))

# 4) 드라이버
sheet("시장 드라이버 (증거 강도순)", "드라이버", ["순위", "요인", "근거", "가중치"], [7, 32, 95, 16],
      d["drivers"] + [["통화정책", "한은·Fed 기대", d["policy"], ""]])

# 5) 업종 (확보된 업종만, 내림차순 정렬 후 RANK)
sec = sorted(d["sectors"], key=lambda x: -x[1])
last_sec = 3 + len(sec)
rows = [[nm, v, f"=RANK(B{i},$B$4:$B${last_sec})", f"=RANK(B{i},$B$4:$B${last_sec},1)", d.get("sector_notes", {}).get(nm, "")]
        for i, (nm, v) in enumerate(sec, 4)]
ws, _ = sheet(d.get("sector_title", "KOSPI 업종 등락률"), "업종", ["업종", CHG, "상승순위", "하락순위", "비고"], [24, 12, 10, 10, 30], rows,
              fmts={2: PCT, 3: '0', 4: '0'}, note=d.get("sector_note", SRC))
ws["G3"] = "해석"; ws["G3"].font = HF; ws["G3"].fill = HFILL
ws["G4"] = d["sector_comment"]; ws["G4"].alignment = WRAP; ws["G4"].font = BLACK
ws.column_dimensions["G"].width = 60; ws.row_dimensions[4].height = 90

# 6) 종목
sheet("주요 종목", "종목", ["기업명(종목코드)", "종가(원)", CHG, "원인", "헤지펀드 시각", "근거"], [28, 13, 10, 50, 20, 58],
      d["stocks"], fmts={2: '#,##0', 3: PCT}, note="빈 종가 = 출처에 종가 미보도. " + d.get("stock_note", ""))

# 7) 실적·이벤트
sheet(d.get("events_title", "실적·공시·지표 리뷰"), "실적이벤트", ["항목", "내용", "주가/시장 반응", "시사점"], [32, 65, 18, 45], d["events"])

# 8) 일정
sheet(d.get("today_sheet_title", "내일 볼 것 (한국시각)"), "일정", ["시점", "이벤트", "Implied move(±)", "포인트"], [16, 55, 16, 48], d["today"],
      fmts={3: '0.0%'})
if d.get("weekend"):
    sheet(d.get("weekend_sheet_title", "주말 뉴스 (금요일 장 마감 후)"), "주말뉴스", ["시점", "뉴스", "시장 영향"], [18, 70, 60], d["weekend"])
sheet("포지셔닝 코멘트 (절대수익 관점)", "포지셔닝", ["구분", "내용"], [16, 115], d["positioning"], note=None)
ws, _ = sheet("Sources", "Sources", ["출처", "URL"], [42, 120], d["sources"], note=d.get("data_note"))
for r in range(4, 4 + len(d["sources"])):
    c = ws.cell(row=r, column=2); c.hyperlink = c.value
    c.font = Font(name=FONT, color="0563C1", underline="single", size=10)

# 0) 요약 (맨 앞)
ws = wb.active; ws.title = "요약"
ws["A1"] = f"{TITLE_TXT} | {d['date']}"; ws["A1"].font = TITLE
ws["A2"] = f"작성 {d['written']} {d.get('written_kst', '')} KST · 범례: 파란 글씨 = 출처 입력값, 검정 = 수식, 초록 = 타 시트 링크"; ws["A2"].font = SUB
items = [("Headline", d["headline_bold"] + " " + d["headline_rest"], None, BLACK)]
if "KOSPI" in row_of and "KOSDAQ" in row_of:
    items.append(("KOSPI − KOSDAQ 상대성과 (%p)", f"={ref('KOSPI','C')}-{ref('KOSDAQ','C')}", '+0.00%;-0.00%', BLACK))
if "국고채 3년" in row_of and "국고채 10년" in row_of:
    items.append(("국고채 3s10s 스프레드 (bp)", f"=({ref('국고채 10년','B')}-{ref('국고채 3년','B')})*10000", '0.0', BLACK))
if "원/달러" in row_of:
    items.append(("원/달러", f"={ref('원/달러','B')}", '#,##0.0"원"', GREEN))
for m in markets:
    if "외국인" in investors:
        items.append((f"{m} 외국인 순매수 (억원)", f"={fref('외국인', m)}", EOK, GREEN))
if "개인" in investors and "외국인" in investors and markets:
    m0 = markets[0]
    items.append((f"{m0} 외국인 vs 개인 (억원, 외국인−개인)", f"={fref('외국인', m0)}-{fref('개인', m0)}", EOK, BLACK))
if d.get("breadth"):
    items.append(("KOSPI 상승/하락 종목 비율", f"='지수·매크로'!B{br+1}/'지수·매크로'!C{br+1}", '0.00"x"', BLACK))
items.append((f"업종 상승 개수 ({len(sec)}개 중)", f"=COUNTIF('업종'!B4:B{last_sec},\">0\")", '0', BLACK))
if d.get("path"):
    nrow = 3 + len(d["path"]["rows"])
    items.append(("주간 최저 종가 (KOSPI)", f"=MIN('일별경로'!B4:B{nrow})", '#,##0.00', BLACK))
    items.append(("주간 일간 최대 낙폭", f"=MIN('일별경로'!C4:C{nrow})", PCT, BLACK))
items.append(("포지셔닝 판단", d["positioning"][0][1], None, BLACK))
for i, (k, v, f, font) in enumerate(items, 4):
    a = ws.cell(row=i, column=1, value=k); a.font = BOLD; a.border = BD; a.alignment = WRAP
    c = ws.cell(row=i, column=2, value=v); c.font = font; c.border = BD; c.alignment = WRAP
    if f: c.number_format = f
ws.column_dimensions["A"].width = 38; ws.column_dimensions["B"].width = 95
ws.row_dimensions[4].height = 62; ws.row_dimensions[3 + len(items)].height = 48
wb.save(out); print(out)
