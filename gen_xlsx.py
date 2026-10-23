"""미국증시 브리핑 엑셀 생성기 v2 (데일리·위클리 공용). 사용: python gen_xlsx.py data.json out.xlsx
v2: data.json의 "scores_from"(score.py 출력 경로)이 있으면 스코어 시트·요약 행을 자동 병합."""
import json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

d = json.load(open(sys.argv[1], encoding="utf-8")); out = sys.argv[2]
if d.get("scores_from"):
    p = d["scores_from"]
    if not os.path.exists(p):
        p = os.path.join(os.path.dirname(os.path.abspath(sys.argv[1])), d["scores_from"])
    s = json.load(open(p, encoding="utf-8"))
    d.setdefault("scores_detail", s["scores_detail"]); d.setdefault("scores_title", s["scores_title"])
    d.setdefault("scores_note", s.get("scores_note", ""))
    d.setdefault("doc", {}).setdefault("scores", s["doc_scores"])
FONT = "Malgun Gothic"
BLUE = Font(name=FONT, color="0000FF", size=10); BLACK = Font(name=FONT, size=10)
BOLD = Font(name=FONT, bold=True, size=10); SUB = Font(name=FONT, italic=True, size=9, color="595959")
HF = Font(name=FONT, bold=True, color="FFFFFF", size=10); HFILL = PatternFill("solid", fgColor="1F3864")
TITLE = Font(name=FONT, bold=True, size=14)
b = Side(style="thin", color="BFBFBF"); BD = Border(left=b, right=b, top=b, bottom=b)
WRAP = Alignment(wrap_text=True, vertical="top")
PCT = '+0.00%;-0.00%;0.00%'; BP = '+0"bp";-0"bp";0"bp"'
TITLE_TXT = d.get("title", "미국증시 데일리 브리핑")
BASIS = f"기준: {d['date']} {d.get('basis_label', '정규장 마감 기준')} / 작성 {d['written']}"
CHG = d.get("chg_label", "등락률")
SRC = "Source: " + ", ".join(s[0] for s in d["sources"][:5]) + " (Sources 시트 참조)"
wb = Workbook()

def sheet(title, name, headers, widths, rows, fmts=None, note=SRC):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE; ws["A2"] = BASIS; ws["A2"].font = SUB
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=3, column=i, value=h); c.font = HF; c.fill = HFILL; c.border = BD
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, 4):
        for ci, v in enumerate(row, 1):
            c = ws.cell(row=r, column=ci, value=v); c.border = BD; c.alignment = WRAP
            num = isinstance(v, (int, float)) and not isinstance(v, bool)
            c.font = BLUE if num else BLACK
            if fmts and ci in fmts and (num or (isinstance(v, str) and v.startswith("="))):
                c.number_format = fmts[ci]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    end = 4 + len(rows)
    if note: ws.cell(row=end, column=1, value=note).font = SUB
    return ws, end

idx = d["indices"]
rows = [[x[0], x[1], x[2], x[3], x[4]] for x in idx]
ws, end = sheet("지수·매크로", "지수·매크로", ["항목", "종가/수준", CHG, "금리변동", "비고"], [30, 14, 11, 11, 44], rows,
                fmts={2: '#,##0.00', 3: PCT, 4: BP})
row_of = {}
for i, x in enumerate(idx, 4):
    row_of[x[0]] = i
    if len(x) > 5 and x[5]: ws.cell(row=i, column=2).number_format = '0.00%'
br = end + 2
if d.get("breadth"):
    ws.cell(row=br, column=1, value="NYSE Breadth").font = BOLD
    ws.cell(row=br, column=2, value="상승").font = BOLD; ws.cell(row=br, column=3, value="하락").font = BOLD
    ws.cell(row=br + 1, column=1, value="종목 수").font = BLACK
    ws.cell(row=br + 1, column=2, value=d["breadth"][0]).font = BLUE; ws.cell(row=br + 1, column=3, value=d["breadth"][1]).font = BLUE

def ref(name, col):
    return f"'지수·매크로'!{col}{row_of[name]}"

# 바닥·고점 신호 스코어 산출 내역: [구분, 영역/지표, 값, 기준, 점수(숫자|null), 비고]
if d.get("scores_detail"):
    sheet(d.get("scores_title", "바닥·고점 신호 스코어 산출 내역"), "스코어", ["구분", "영역/지표", "값", "기준", "점수", "비고"],
          [20, 30, 26, 36, 9, 46], d["scores_detail"], fmts={5: '0'}, note=d.get("scores_note", SRC))

sheet("시장 드라이버 (증거 강도순)", "드라이버", ["순위", "요인", "근거", "가중치"], [7, 32, 90, 18],
      d["drivers"] + [["Fed", "Fed 기대", d["fed"], ""]])

sec = sorted(d["sectors"], key=lambda x: -x[1])
rows = [[n, v, f"=RANK(B{i},$B$4:$B${3+len(sec)})", f"=RANK(B{i},$B$4:$B${3+len(sec)},1)", d.get("sector_notes", {}).get(n, "")]
        for i, (n, v) in enumerate(sec, 4)]
ws, _ = sheet("S&P 500 섹터 등락률", "섹터", ["섹터", CHG, "상승순위", "하락순위", "비고"], [26, 12, 10, 10, 30], rows,
              fmts={2: PCT, 3: '0', 4: '0'})
ws["G3"] = "해석"; ws["G3"].font = HF; ws["G3"].fill = HFILL
ws["G4"] = d["sector_comment"]; ws["G4"].alignment = WRAP; ws["G4"].font = BLACK
ws.column_dimensions["G"].width = 60; ws.row_dimensions[4].height = 75
last_sec = 3 + len(sec)

ws, _ = sheet("주요 종목", "종목", ["기업명(티커)", "종가($)", CHG, "원인", "헤지펀드 시각", "근거"], [30, 11, 10, 50, 22, 55],
              d["stocks"], fmts={2: '#,##0.00', 3: PCT},
              note="빈 종가 = 출처에 종가 미보도. " + d.get("stock_note", ""))

sheet(d.get("earnings_sheet_title", f"실적 리뷰 ({d['date']} 발표)"), "실적리뷰", ["기업명(티커)", "결과", "주가 반응", "시사점"], [32, 65, 18, 45], d["earnings"])
sheet(d.get("today_sheet_title", "오늘 볼 것 (미국시간)"), "오늘일정", ["시간", "이벤트", "Implied move(±)", "포인트"], [10, 55, 16, 45], d["today"],
      fmts={3: '0.0%'})
if d.get("weekend"):
    sheet(d.get("weekend_sheet_title", "주말 뉴스 (미국 금요일 장 마감 후)"), d.get("weekend_sheet_name", "주말뉴스"), ["시점", "뉴스", "시장 영향"], [16, 70, 60], d["weekend"])
sheet("포지셔닝 코멘트 (절대수익 관점)", "포지셔닝", ["구분", "내용"], [16, 110], d["positioning"], note=None)
ws, _ = sheet("Sources", "Sources", ["출처", "URL"], [40, 120], d["sources"], note=d.get("data_note"))
for r in range(4, 4 + len(d["sources"])):
    c = ws.cell(row=r, column=2); c.hyperlink = c.value
    c.font = Font(name=FONT, color="0563C1", underline="single", size=10)

ws = wb.active; ws.title = "요약"
ws["A1"] = f"{TITLE_TXT} | {d['date']}"; ws["A1"].font = TITLE
ws["A2"] = f"작성 {d['written']} · 범례: 파란 글씨 = 출처 입력값, 검정 = 수식"; ws["A2"].font = SUB
items = [("Headline", d["headline_bold"] + " " + d["headline_rest"], None)]
for s in d.get("doc", {}).get("scores", []):  # [구분, 점수, 판정, 국면/유형, 핵심 근거, 전환 조건]
    items.append((s[0], f"{s[1]} · {s[2]} · {s[3]} | 근거: {s[4]} | 전환 조건: {s[5]}", None))
if "Russell 2000" in row_of and "S&P 500" in row_of:
    items.append(("Russell 2000 − S&P 500 (%p)", f"={ref('Russell 2000','C')}-{ref('S&P 500','C')}", '+0.00%;-0.00%'))
if "S&P 500 Equal Weight (RSP)" in row_of:
    items.append(("S&P 동일가중(RSP) − S&P 500 (%p)", f"={ref('S&P 500 Equal Weight (RSP)','C')}-{ref('S&P 500','C')}", '+0.00%;-0.00%'))
if "UST 2Y" in row_of and "UST 10Y" in row_of:
    items.append(("2s10s 스프레드 (bp)", f"=({ref('UST 10Y','B')}-{ref('UST 2Y','B')})*10000", '0.0'))
if d.get("breadth"):
    items.append(("NYSE 상승/하락 비율", f"='지수·매크로'!B{br+1}/'지수·매크로'!C{br+1}", '0.00"x"'))
items.append((f"섹터 상승 개수 ({len(sec)}개 중)", f"=COUNTIF('섹터'!B4:B{last_sec},\">0\")", '0'))
items.append(("포지셔닝 판단", d["positioning"][0][1], None))
for i, (k, v, f) in enumerate(items, 4):
    a = ws.cell(row=i, column=1, value=k); a.font = BOLD; a.border = BD; a.alignment = WRAP
    c = ws.cell(row=i, column=2, value=v); c.font = BLACK; c.border = BD; c.alignment = WRAP
    if f: c.number_format = f
ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 95
ws.row_dimensions[4].height = 48; ws.row_dimensions[3 + len(items)].height = 48
wb.save(out); print(out)
